#!/usr/bin/env python3
"""
FORMAT QA GATE - uapf-quality-gates (BOTH ENGINES, MANDATORY)
=============================================================
Mechanical linter for Genie's formatting laws. Run on every chapter DOCX
(and optionally its rendered PDF) BEFORE the preview ships. FAIL blocks
advancement: fix and re-run. Codex sessions run this exactly like Claude
sessions; "the instructions said so" is replaced by "the gate said PASS".

Checks:
  F1 em dashes (banned everywhere)
  F2 tiny type: body runs below the floor (12pt default / 13pt senior;
     9.5pt hard floor for captions, folios, and meta lines)
  F3 8.5x11 heading law: Heading-1 chapter titles must be >= 28pt
  F4 hex color codes presented in body text (fabricated materials tell)
  F5 template stamping: identical sentences (8+ words) repeated 3+ times
  F6 publisher name / ISBN placeholder in reader-facing text
  F7 page density (needs --pdf): unit pages whose content fills under
     ~55% of the live area are listed; more than 15% of body pages
     under-filled = FAIL (openers and section ends excluded by tolerance)

Usage:
  python format_qa.py <chapter_or_book.docx> [--pdf rendered.pdf]
                      [--senior] [--trim 8.5x11] [--json out.json]
Exit 0 PASS / 1 WARN / 2 FAIL.
"""
import argparse, json, re, sys
from collections import Counter

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--pdf")
    ap.add_argument("--senior", action="store_true")
    ap.add_argument("--trim", default="8.5x11")
    ap.add_argument("--json")
    ap.add_argument("--require-images", dest="require_images", type=int, default=None,
                    help="FAIL unless the chapter embeds at least N real images "
                         "(compulsory inline generation law, operator directive 2026-08-19)")
    ap.add_argument("--min-para-lines", dest="min_para_lines", type=int, default=None,
                    help="PARAGRAPH LAW: FAIL when body paragraphs are choppy; pass 10 "
                         "for prose and textbook chapters (operator directive 2026-08-30)")
    a = ap.parse_args()

    import docx as docxlib
    d = docxlib.Document(a.docx)
    fails, warns = [], []

    # F0 compulsory inline images: a chapter whose plan calls for images may not
    # ship its preview imageless; images are generated DURING drafting, never after.
    if a.require_images is not None:
        import re as _re
        with open(a.docx, "rb") as _f:
            pass
        import zipfile as _zf
        _z = _zf.ZipFile(a.docx)
        _xml = _z.read("word/document.xml").decode("utf-8", "ignore")
        n_img = _xml.count("<w:drawing>") + _xml.count("<w:drawing ")
        media = [n for n in _z.namelist() if n.startswith("word/media/")]
        placeholders = len(_re.findall(r"(?i)image placeholder|\[image|IMG-R?\d+\s*placeholder", _xml))
        if n_img < a.require_images or not media:
            fails.append(f"F0: chapter embeds {n_img} image(s) with {len(media)} media file(s); "
                         f"the plan requires at least {a.require_images}. Images are generated "
                         f"DURING the chapter (compulsory), never deferred.")
        if placeholders:
            fails.append(f"F0: {placeholders} image PLACEHOLDER(s) found; placeholders never ship, "
                         f"the real image is generated inline before the preview.")
    body_floor = 13.0 if a.senior else 12.0
    hard_floor = 9.5
    big_trim = a.trim.lower().startswith("8.5")

    # F1 em dash
    em = sum(p.text.count("—") for p in d.paragraphs)
    if em:
        fails.append(f"F1: {em} em dash(es) found (banned everywhere)")

    # F2 type floors + F3 heading law
    tiny, small_body, h1_small = [], [], []
    for p in d.paragraphs:
        style = (p.style.name if p.style else "") or ""
        for r in p.runs:
            if not r.text.strip():
                continue
            sz = r.font.size.pt if r.font.size else None
            if sz is None:
                continue
            if sz < hard_floor:
                tiny.append((sz, r.text[:30]))
            elif sz < body_floor and len(r.text.split()) >= 6 and not style.startswith("Heading"):
                small_body.append((sz, r.text[:30]))
            if big_trim and style == "Heading 1" and sz < 28:
                h1_small.append((sz, p.text[:40]))
    if tiny:
        fails.append(f"F2: {len(tiny)} run(s) below the {hard_floor}pt hard floor, e.g. {tiny[0]}")
    if len(small_body) > 20:
        fails.append(f"F2: {len(small_body)} body-length runs below the {body_floor}pt body floor, e.g. {small_body[0]}")
    elif small_body:
        warns.append(f"F2: {len(small_body)} body-length runs below {body_floor}pt")
    if h1_small:
        fails.append(f"F3: {len(h1_small)} chapter heading(s) under 28pt at 8.5x11, e.g. {h1_small[0]}")

    # F4 hex codes in text
    hexes = [p.text for p in d.paragraphs if re.search(r"#[0-9A-Fa-f]{6}\b", p.text)]
    if hexes:
        fails.append(f"F4: hex color codes in reader-facing text ({len(hexes)} paragraph(s)); "
                     "materials/colors must be real-world specs")

    # F5 template stamping
    sent = Counter()
    for p in d.paragraphs:
        for s in re.split(r"(?<=[.!?])\s+", p.text):
            s = s.strip()
            if len(s.split()) >= 8:
                sent[s.lower()] += 1
    stamped = [(s, c) for s, c in sent.items() if c >= 3]
    if stamped:
        worst = max(stamped, key=lambda x: x[1])
        fails.append(f"F5: {len(stamped)} sentence(s) template-stamped 3+ times "
                     f"(worst repeated {worst[1]}x: \"{worst[0][:50]}...\")")

    # F6 publisher / ISBN placeholder
    low = "\n".join(p.text for p in d.paragraphs).lower()
    if "isbn: [" in low or "isbn [to be" in low or "[isbn" in low:
        fails.append("F6: ISBN placeholder present")
    if "pegasus press" in low:
        fails.append("F6: publisher name in reader-facing text")

    # F8 tab / first-line-indent paragraphs (PARAGRAPH LAW, operator directive
    # 2026-08-30, every niche, both engines): manuscripts use BLOCK paragraphs
    # only. Any body paragraph with a first-line indent, or opening with a tab
    # character, is a production failure.
    indented = []
    for p in d.paragraphs:
        style = (p.style.name if p.style else "") or ""
        if style.startswith("Heading"):
            continue
        fli = p.paragraph_format.first_line_indent
        if (fli is not None and fli and fli.pt > 1) or p.text.startswith("\t"):
            indented.append(p.text[:40])
    if indented:
        fails.append(f"F8: {len(indented)} tab/first-line-indented paragraph(s) "
                     f"(block paragraphs only), e.g. \"{indented[0]}\"")

    # F9 paragraph density (PARAGRAPH LAW, operator directive 2026-08-30):
    # prose and textbook body paragraphs must be FULL, about 10 rendered lines
    # before a break, never choppy 2-3 sentence AI-style fragments. Opt in
    # with --min-para-lines (the formatting standard passes 10 for prose /
    # textbook chapters); list items, headings, captions, and short lead-ins
    # are exempt via the 20-word candidate floor.
    if a.min_para_lines:
        wpl = 14 if big_trim else 10   # approx words per rendered line
        cand = []
        for p in d.paragraphs:
            style = (p.style.name if p.style else "") or ""
            if style.startswith(("Heading", "List")) or "Caption" in style:
                continue
            w = len(p.text.split())
            if w >= 20:
                cand.append(w / wpl)
        if cand:
            short = [round(l, 1) for l in cand if l < a.min_para_lines]
            share = len(short) / len(cand)
            if share > 0.40:
                fails.append(
                    f"F9: choppy paragraphs: {len(short)}/{len(cand)} body paragraphs "
                    f"run under ~{a.min_para_lines} rendered lines ({share:.0%}). Merge and "
                    f"deepen them into full paragraphs (about {a.min_para_lines * wpl}+ words "
                    f"each); never break a thought every 2-3 sentences.")
            elif share > 0.20:
                warns.append(f"F9: {len(short)}/{len(cand)} body paragraphs under "
                             f"~{a.min_para_lines} rendered lines ({share:.0%})")

    # F7 density via rendered PDF
    if a.pdf:
        import fitz
        doc = fitz.open(a.pdf)
        under, body_pages = [], 0
        for pno in range(doc.page_count):
            page = doc[pno]
            pw, ph = page.rect.width, page.rect.height
            live_top, live_bot = ph * 0.08, ph * 0.92
            blocks = [b for b in page.get_text("blocks") if b[4].strip()] + \
                     [[*img["bbox"], ""] for img in page.get_image_info()]
            # exclude running headers/folios: blocks living entirely in the
            # header or footer bands are furniture, not content
            blocks = [b for b in blocks if b[1] < live_bot * 0.96 and b[3] > live_top * 1.1]
            if not blocks:
                continue
            body_pages += 1
            lowest = max(min(b[3], live_bot) for b in blocks)
            fill = (min(lowest, live_bot) - live_top) / (live_bot - live_top)
            if fill < 0.55:
                under.append((pno + 1, round(fill, 2)))
        if body_pages:
            share = len(under) / body_pages
            if share > 0.15:
                fails.append(f"F7: {len(under)}/{body_pages} pages fill under 55% of the live area "
                             f"({share:.0%}); e.g. pages {[p for p, _ in under[:8]]}")
            elif under:
                warns.append(f"F7: {len(under)} under-filled page(s): {[p for p, _ in under[:6]]}")

    verdict = "FAIL" if fails else ("WARN" if warns else "PASS")
    print(f"{verdict}  format_qa on {a.docx}")
    for f in fails: print("  FAIL:", f)
    for w in warns: print("  warn:", w)
    if a.json:
        json.dump({"verdict": verdict, "fails": fails, "warns": warns},
                  open(a.json, "w", encoding="utf-8"), indent=1)
    sys.exit(0 if verdict == "PASS" else (1 if verdict == "WARN" else 2))

if __name__ == "__main__":
    main()
