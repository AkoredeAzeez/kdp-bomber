#!/usr/bin/env python3
"""
TEXT HYGIENE - uapf-quality-gates (rendering cleanliness, VISIBLE QA step)
==========================================================================
Removes genuinely STRAY invisible / control characters that break DOCX and
EPUB rendering or trip KDP ingestion, and reports exactly what it did. This is
file cleanliness, run and logged like any other gate.

  python text_hygiene.py <file.docx>            # inspect only (report, exit 1 if issues)
  python text_hygiene.py <file.docx> --fix      # clean in place (backup kept), report

SCOPE (deliberate and narrow):
  REMOVED  - zero-width space U+200B, word joiner U+2060, BOM/ZWNBSP U+FEFF
             mid-text, soft hyphen U+00AD (renders as a phantom hyphen), and
             C0/C1 control chars except tab/newline/carriage-return.
  FLAGGED (not touched) - U+FFFD replacement char (signals real encoding data
             loss, needs human review), and ZWJ/ZWNJ U+200D/U+200C (legitimate
             in Arabic, Indic, and emoji; never blind-stripped in a
             multi-language studio).
  PRESERVED - non-breaking space, normal hyphens, all real letters and
             punctuation, and every legitimate typographic mark.

NOT IN SCOPE, BY DESIGN: this does NOT target, detect, or defeat any AI
provenance or watermark system (SynthID, C2PA, vendor marks, statistical text
watermarks) and does NOT alter authorship or AI-disclosure status. It is a
rendering-cleanliness pass, not an anti-detection tool. AI-disclosure stays
truthful in the publish manifest.
"""
import argparse, os, shutil, sys
from collections import Counter

REMOVE = {
    "​": "zero-width space",
    "⁠": "word joiner",
    "﻿": "BOM / zero-width no-break space",
    "­": "soft hyphen (phantom hyphen)",
}
# C0 (0x00-0x1F) and C1 (0x80-0x9F) controls except \t \n \r
REMOVE_CTRL = {chr(c) for c in list(range(0x00, 0x20)) + list(range(0x80, 0xA0))} - {"\t", "\n", "\r"}
FLAG = {
    "�": "replacement char (encoding data loss - review the source)",
    "‍": "zero-width joiner (legitimate in some scripts/emoji - left intact)",
    "‌": "zero-width non-joiner (legitimate in some scripts - left intact)",
}

def scan_text(t):
    removed, flagged = Counter(), Counter()
    for ch in t:
        if ch in REMOVE or ch in REMOVE_CTRL:
            removed[ch] += 1
        elif ch in FLAG:
            flagged[ch] += 1
    return removed, flagged

def clean_text(t):
    out = []
    for ch in t:
        if ch in REMOVE or ch in REMOVE_CTRL:
            continue
        out.append(ch)
    return "".join(out)

def process_docx(path, fix):
    import docx as docxlib
    d = docxlib.Document(path)
    removed, flagged = Counter(), Counter()

    def handle_par(p):
        for r in p.runs:
            rem, flg = scan_text(r.text or "")
            removed.update(rem); flagged.update(flg)
            if fix and rem:
                r.text = clean_text(r.text)

    for p in d.paragraphs:
        handle_par(p)
    for tb in d.tables:
        for row in tb.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    handle_par(p)
    # headers/footers
    for sec in d.sections:
        for hf in (sec.header, sec.footer):
            for p in hf.paragraphs:
                handle_par(p)

    if fix and sum(removed.values()):
        bak = path + ".prehygiene.bak"
        if not os.path.exists(bak):
            shutil.copy2(path, bak)
        d.save(path)
    return removed, flagged

def process_text_file(path, fix):
    raw = open(path, encoding="utf-8-sig", errors="replace").read()
    removed, flagged = scan_text(raw)
    if fix and sum(removed.values()):
        bak = path + ".prehygiene.bak"
        if not os.path.exists(bak):
            shutil.copy2(path, bak)
        open(path, "w", encoding="utf-8").write(clean_text(raw))
    return removed, flagged

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--fix", action="store_true")
    a = ap.parse_args()
    if not os.path.exists(a.file):
        sys.exit("file not found: " + a.file)
    ext = os.path.splitext(a.file)[1].lower()
    if ext == ".docx":
        removed, flagged = process_docx(a.file, a.fix)
    elif ext in (".txt", ".md", ".html", ".xml"):
        removed, flagged = process_text_file(a.file, a.fix)
    else:
        sys.exit("unsupported: send a .docx, .txt, .md, .html, or .xml")

    def show(label, counter, names):
        if not counter:
            print(f"  {label}: none")
            return
        for ch, n in counter.most_common():
            nm = names.get(ch, "control U+%04X" % ord(ch))
            print(f"  {label}: {n:5}  {nm}")

    print(("CLEANED" if a.fix else "INSPECTED"), os.path.basename(a.file))
    show("removed" if a.fix else "would remove", removed, {**REMOVE})
    show("flagged (left intact)", flagged, FLAG)
    total = sum(removed.values())
    if a.fix:
        print(f"Text hygiene: removed {total} stray invisible/control char(s). "
              f"Backup: {os.path.basename(a.file)}.prehygiene.bak" if total else
              "Text hygiene: clean, nothing to remove.")
        sys.exit(0)
    else:
        print(f"Text hygiene: {total} stray char(s) found."
              if total else "Text hygiene: clean.")
        sys.exit(1 if total else 0)

if __name__ == "__main__":
    main()
