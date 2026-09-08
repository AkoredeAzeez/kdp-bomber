#!/usr/bin/env python3
"""
CMS RESULT DOCX - uapf-close-title-strategist (OPERATOR-ONLY)
=============================================================
Builds the operator-facing DOCX deliverable for every CMS run: a
two-column comparison table per niche:

  ORIGINAL TITLE AND SUBTITLE  |  NEW TITLE AND SUBTITLE

Input JSON:
{
  "title": "CMS Sweep Results",
  "sections": [
    { "heading": "Crime Puzzles and Activity Books (Amazon.com, 2026-08-16)",
      "rows": [
        { "orig_title": "...", "orig_subtitle": "...", "meta": "ASIN ... | BSR ... | rating",
          "new_title": "...", "new_subtitle": "..." }, ... ] } ]
}

Usage: python cms_docx.py <data.json> <out.docx>
"""
import json, os, sys
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PRINCIPAL = RGBColor(0x1F, 0x4E, 0x5F)
ACCENT = RGBColor(0xE5, 0x8E, 0x3A)
GRAY = RGBColor(0x5A, 0x6A, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:fill"), hexcolor); tcPr.append(sh)


def build(data, out):
    d = docx.Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(0.7)
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(9.5)

    t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(data.get("title", "CMS RESULTS")); r.bold = True
    r.font.size = Pt(20); r.font.color.rgb = PRINCIPAL

    for sec in data.get("sections", []):
        h = d.add_paragraph(); r = h.add_run(sec.get("heading", ""))
        r.bold = True; r.font.size = Pt(14); r.font.color.rgb = ACCENT
        h.paragraph_format.space_before = Pt(12)

        rows = sec.get("rows", [])
        tb = d.add_table(rows=1 + len(rows), cols=2)
        tb.style = "Table Grid"
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = tb.rows[0].cells
        for ci, htxt in enumerate(("ORIGINAL TITLE AND SUBTITLE", "NEW TITLE AND SUBTITLE")):
            p = hdr[ci].paragraphs[0]
            r = p.add_run(htxt); r.bold = True; r.font.color.rgb = WHITE; r.font.size = Pt(10)
            shade(hdr[ci], "1F4E5F")
        for ri, row in enumerate(rows, start=1):
            c0, c1 = tb.rows[ri].cells
            p = c0.paragraphs[0]
            r = p.add_run(row.get("orig_title", "")); r.bold = True; r.font.size = Pt(9.5)
            if row.get("orig_subtitle"):
                p2 = c0.add_paragraph()
                p2.add_run(row["orig_subtitle"]).font.size = Pt(9)
            if row.get("meta"):
                p3 = c0.add_paragraph()
                r = p3.add_run(row["meta"]); r.italic = True; r.font.size = Pt(8); r.font.color.rgb = GRAY
            p = c1.paragraphs[0]
            r = p.add_run(row.get("new_title", "")); r.bold = True
            r.font.size = Pt(10); r.font.color.rgb = PRINCIPAL
            if row.get("new_subtitle"):
                p2 = c1.add_paragraph()
                p2.add_run(row["new_subtitle"]).font.size = Pt(9)
            if ri % 2 == 0:
                shade(c0, "F2F7F8"); shade(c1, "F2F7F8")
        for tr in tb.rows:
            tr.cells[0].width = Inches(3.5); tr.cells[1].width = Inches(3.6)
        d.add_paragraph()

    ft = d.add_paragraph(); ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = ft.add_run("CMS candidates only: every selected title must pass the trademark "
                   "gate before production. Search adjacency reflects the research date; "
                   "Amazon results vary and nothing here guarantees visibility or ranking.")
    r.italic = True; r.font.size = Pt(8); r.font.color.rgb = GRAY
    d.save(out)
    print(f"CMS DOCX: {out} ({os.path.getsize(out)//1024} KB, "
          f"{sum(len(s.get('rows', [])) for s in data.get('sections', []))} comparisons)")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    build(json.load(open(sys.argv[1], encoding="utf-8")), sys.argv[2])
