#!/usr/bin/env python3
"""
Genie - Production Dashboard generator (professional).
Shows only genuine Genie books (a book_lock.md production record and/or a catalog.json entry),
attributes each to the agent(s) that produced it - Claude Code and/or Codex - and writes a
self-contained Genie_Dashboard.html.

Portable: locates your catalog + project folders automatically, so it works on the operator
machine and on any client install. Re-run any time to refresh:  python genie_dashboard.py
"""
import os, re, json, html, glob, datetime, collections

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
HOME = os.path.expanduser("~")
OUT = os.path.join(SCRIPT_DIR, "Genie_Dashboard.html")

CANDIDATE_CATALOGS = [
    os.path.join(SCRIPT_DIR, "catalog.json"),
    os.path.join(SCRIPT_DIR, "catalog", "catalog.json"),
    os.path.join(SCRIPT_DIR, "..", "catalog.json"),
    os.path.join(HOME, "Genie", "catalog.json"),
    os.path.join(HOME, "Documents", "UAPF_Projects", "catalog.json"),
]
PROJECT_ROOTS = [
    os.path.join(HOME, "Genie", "projects"),
    os.path.join(HOME, "Documents", "UAPF_Projects"),
    os.path.join(SCRIPT_DIR, "projects"),
    SCRIPT_DIR,
]
# Folder-name fragments that are demos/tests, never real books. Kept generic on purpose:
# real books are discovered only via a book_lock.md or a catalog entry, so other folders
# are never picked up regardless.
IGNORE = ["sample", "demo", "__pycache__", ".git", "node_modules", "_backup"]

STATUS_MAP = {"finished": "finished", "complete": "finished", "completed": "finished", "done": "finished",
              "published": "published", "live": "published",
              "in progress": "in progress", "in-progress": "in progress", "wip": "in progress",
              "drafting": "in progress", "draft": "in progress"}
STATUS_COLORS = {"published": "#2563eb", "finished": "#16a34a", "in progress": "#d97706", "unknown": "#64748b"}
AGENT_COLORS = {"Claude Code": "#d97706", "Codex": "#0ea5e9", "Both": "#7c3aed", "-": "#64748b"}
TM_COLORS = {"CLEARED": "#16a34a", "CLEAR": "#16a34a", "RISK": "#dc2626", "UNVERIFIED": "#64748b", "-": "#64748b"}

def norm_status(s): return STATUS_MAP.get((s or "").strip().lower(), (s or "unknown").strip().lower())
def ignored(name):
    n = (name or "").lower()
    return any(frag in n for frag in IGNORE)

def find_catalog():
    for p in CANDIDATE_CATALOGS:
        if os.path.isfile(p):
            try:
                with open(p, encoding="utf-8") as f: return p, json.load(f)
            except Exception: pass
    return None, []

def find_booklocks():
    seen, out = set(), []
    for root in PROJECT_ROOTS:
        if not os.path.isdir(root): continue
        # depth-limited: root/*/book_lock.md and root/book_lock.md
        for pat in (os.path.join(root, "book_lock.md"), os.path.join(root, "*", "book_lock.md")):
            for bl in glob.glob(pat):
                folder = os.path.dirname(bl)
                rp = os.path.normcase(os.path.abspath(folder))
                if rp in seen or ignored(os.path.basename(folder)): continue
                # the install root ships a sample book_lock.md (template); a folder that
                # itself holds a projects/ directory is an install root, never a book
                if rp == os.path.normcase(os.path.abspath(SCRIPT_DIR)) or                    os.path.isdir(os.path.join(folder, "projects")): continue
                seen.add(rp); out.append(folder)
    return out

def parse_booklock(folder):
    bl = os.path.join(folder, "book_lock.md")
    d = {}
    if not os.path.isfile(bl): return d
    try: t = open(bl, encoding="utf-8", errors="replace").read()
    except Exception: return d
    m = re.search(r"Working title:\s*(.+)", t)
    if m: d["title"] = m.group(1).strip()
    m = re.search(r"selected_skill=uapf-niche-([a-z0-9\-]+)", t)
    if m: d["niche"] = m.group(1).replace("-", " ").title()
    else:
        m = re.search(r"Profile:\s*([A-Z0-9\-]+)", t)
        if m: d["niche"] = m.group(1)
    m = re.search(r"framework=([^;\n]+)", t)
    if m: d["framework"] = m.group(1).strip()
    m = re.search(r"\bVERDICT:\s*(GREEN|YELLOW|RED|CLEAR(?:ED)?|RISK)", t, re.I) or \
        re.search(r"\btrademark_status[=:\s]+(CLEARED|CLEAR|RISK|UNVERIFIED)", t, re.I) or \
        re.search(r"Verdict:\s*(CLEAR(?:ED)?|RISK)", t, re.I)
    if m:
        v = m.group(1).upper()
        d["trademark"] = {"GREEN": "CLEAR", "YELLOW": "RISK", "RED": "RISK"}.get(v, v)
    d["locked"] = "classification_locked=true" in t
    return d

def parse_agents(folder):
    log = os.path.join(folder, "state", "audit_log.jsonl")
    claude = codex = 0
    if os.path.isfile(log):
        try:
            for line in open(log, encoding="utf-8", errors="replace"):
                line = line.strip()
                if not line: continue
                try: ev = json.loads(line)
                except Exception: continue
                e = (ev.get("event") or "")
                if e.startswith("CODEX_TASK_COMPLETED"): codex += 1
                elif e.startswith("CLAUDE_TASK_COMPLETED"): claude += 1
        except Exception: pass
    if claude and codex: label = "Both"
    elif codex: label = "Codex"
    elif claude: label = "Claude Code"
    else: label = None  # default assigned later
    return {"agent": label, "claude_units": claude, "codex_units": codex}

def scan_folder(path):
    docx, cover, latest = [], False, 0.0
    if path and os.path.isdir(path):
        for r, _, files in os.walk(path):
            for fn in files:
                low = fn.lower(); fp = os.path.join(r, fn)
                try: latest = max(latest, os.stat(fp).st_mtime)
                except OSError: pass
                if low.endswith(".docx") and not low.startswith("~$"): docx.append(fp)
                if "cover" in low and low.endswith((".png", ".jpg", ".jpeg", ".pdf")): cover = True
    words = None
    if docx:
        big = max(docx, key=lambda p: os.path.getsize(p))
        try:
            from docx import Document
            doc = Document(big)
            words = sum(len(p.text.split()) for p in doc.paragraphs)
            for tb in doc.tables:
                for row in tb.rows:
                    for c in row.cells: words += len(c.text.split())
        except Exception: words = None
    return {"docx": len(docx), "cover": cover, "latest": latest, "words": words}

def build():
    cat_path, cat = find_catalog()
    books, by_path = [], {}
    def key(p): return os.path.normcase(os.path.abspath(p)) if p else None
    # 1) catalog entries
    for r in cat:
        p = r.get("path", "")
        if ignored(r.get("folder") or os.path.basename(p or "")): continue
        b = {"title": r.get("title"), "author": r.get("author"), "subtitle": r.get("subtitle"),
             "category": r.get("category"), "status": r.get("status"), "date": r.get("date"),
             "path": p, "cover_meta": bool(r.get("cover"))}
        books.append(b); by_path[key(p)] = b
    # 2) book_lock folders not already present
    for folder in find_booklocks():
        if key(folder) in by_path: continue
        books.append({"title": None, "author": None, "subtitle": None, "category": None,
                      "status": None, "date": None, "path": folder, "cover_meta": False})
        by_path[key(folder)] = books[-1]
    # 3) enrich
    enriched = []
    for b in books:
        p = b["path"]
        bl = parse_booklock(p) if p else {}
        ag = parse_agents(p) if p else {"agent": None, "claude_units": 0, "codex_units": 0}
        fs = scan_folder(p) if p else {"docx": 0, "cover": False, "latest": 0, "words": None}
        title = b["title"] or bl.get("title") or (os.path.basename(p) if p else "Untitled")
        # date fallback from folder mtime
        date = b["date"]
        if not date and fs["latest"]:
            date = datetime.datetime.fromtimestamp(fs["latest"]).strftime("%Y-%m-%d")
        enriched.append({
            "title": title, "author": b["author"], "subtitle": b["subtitle"],
            "niche": bl.get("niche") or b["category"] or "-",
            "category": b["category"] or bl.get("niche") or "Uncategorized",
            "framework": bl.get("framework", "-"),
            "trademark": bl.get("trademark", "-"),
            "status_norm": norm_status(b["status"]) if b["status"] else ("finished" if (b["cover_meta"] or fs["cover"]) else "in progress"),
            "agent": ag["agent"] or "Claude Code",
            "claude_units": ag["claude_units"], "codex_units": ag["codex_units"],
            "cover": b["cover_meta"] or fs["cover"], "docx": fs["docx"], "words": fs["words"],
            "date": date or "-", "latest": fs["latest"], "path": p,
        })
    return cat_path, enriched

def e(s): return html.escape(str(s if s is not None else ""))

def bars(counts, palette):
    total = sum(counts.values()) or 1
    out = []
    for k, v in sorted(counts.items(), key=lambda kv: -kv[1]):
        pct = round(100 * v / total)
        col = palette.get(k, "var(--accent)")
        out.append(f'<div class="brow"><div class="blabel" title="{e(k)}">{e(k)}</div>'
                   f'<div class="btrack"><div class="bfill" style="width:{max(pct,4)}%;background:{col}"></div></div>'
                   f'<div class="bval">{v}</div></div>')
    return "\n".join(out) or '<div class="muted">No data</div>'

def pill(text, color): return f'<span class="pill" style="--c:{color}">{e(text)}</span>'

def render(cat_path, books):
    now = datetime.datetime.now().strftime("%d %b %Y, %H:%M")
    total = len(books)
    st = collections.Counter(b["status_norm"] for b in books)
    ni = collections.Counter((b["niche"] if b["niche"] != "-" else "Unclassified") for b in books)
    ag = collections.Counter(b["agent"] for b in books)
    mo = collections.Counter((str(b["date"])[:7] if b["date"] and b["date"] != "-" else "unknown") for b in books)
    manuscripts = sum(b["docx"] for b in books)
    covers = sum(1 for b in books if b["cover"])
    words = sum(b["words"] or 0 for b in books)
    codex_books = sum(1 for b in books if b["agent"] in ("Codex", "Both"))

    def kpi(v, label, sub=""):
        return (f'<div class="kpi"><div class="kval">{e(v)}</div><div class="klabel">{e(label)}</div>'
                + (f'<div class="ksub">{e(sub)}</div>' if sub else "") + "</div>")
    kpis = "".join([
        kpi(total, "Books"),
        kpi(st.get("finished", 0) + st.get("published", 0), "Finished"),
        kpi(st.get("in progress", 0), "In progress"),
        kpi(f"{words:,}", "Words", "in manuscripts"),
        kpi(manuscripts, "Manuscripts", ".docx"),
        kpi(covers, "Covers"),
        kpi(codex_books, "Codex-assisted", "failover used"),
    ])

    rows = []
    for b in sorted(books, key=lambda x: x["latest"], reverse=True):
        stt = b["status_norm"]
        last = datetime.datetime.fromtimestamp(b["latest"]).strftime("%Y-%m-%d") if b["latest"] else "-"
        agsub = ""
        if b["claude_units"] or b["codex_units"]:
            agsub = f'<div class="bsub">CC {b["claude_units"]} · Cx {b["codex_units"]}</div>'
        rows.append(f"""<tr>
          <td><div class="btitle">{e(b['title'])}</div>{f'<div class="bsub">{e((b['subtitle'] or '')[:88])}</div>' if b['subtitle'] else ''}</td>
          <td>{e(b['author'] or '—')}</td>
          <td>{e(b['niche'])}</td>
          <td>{pill(b['agent'], AGENT_COLORS.get(b['agent'], '#64748b'))}{agsub}</td>
          <td>{pill(stt, STATUS_COLORS.get(stt, '#64748b'))}</td>
          <td>{pill(b['trademark'], TM_COLORS.get(b['trademark'], '#64748b')) if b['trademark'] != '-' else '<span class="muted">—</span>'}</td>
          <td class="num">{'✓' if b['cover'] else '—'}</td>
          <td class="num">{f"{b['words']:,}" if b['words'] else '—'}</td>
          <td class="num">{e(b['date'])}</td>
          <td class="num">{last}</td>
        </tr>""")
    table = "\n".join(rows) or '<tr><td colspan="10" class="muted" style="text-align:center;padding:30px">No Genie books found yet. Start one and it will appear here.</td></tr>'
    src = e(cat_path) if cat_path else "project scan"
    return TEMPLATE.format(now=e(now), total=total, kpis=kpis, table=table,
                           status_bars=bars(st, STATUS_COLORS), niche_bars=bars(ni, {}),
                           agent_bars=bars(ag, AGENT_COLORS), month_bars=bars(mo, {}), src=src)

TEMPLATE = r"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Genie - Production Dashboard</title>
<style>
:root{{--bg:#f4f5f7;--card:#ffffff;--ink:#0f172a;--muted:#64748b;--line:#e6e8ec;--accent:#7c3aed;
 --brand1:#7c3aed;--brand2:#4f46e5;--shadow:0 1px 2px rgba(16,24,40,.06),0 1px 3px rgba(16,24,40,.10)}}
@media (prefers-color-scheme:dark){{:root{{--bg:#0b0e13;--card:#141922;--ink:#e7eaf0;--muted:#8b95a5;--line:#232b36;--accent:#a78bfa;--brand1:#7c3aed;--brand2:#6366f1;--shadow:0 1px 2px rgba(0,0,0,.4)}}}}
:root[data-theme=dark]{{--bg:#0b0e13;--card:#141922;--ink:#e7eaf0;--muted:#8b95a5;--line:#232b36;--accent:#a78bfa}}
:root[data-theme=light]{{--bg:#f4f5f7;--card:#fff;--ink:#0f172a;--muted:#64748b;--line:#e6e8ec;--accent:#7c3aed}}
*{{box-sizing:border-box}}
html,body{{margin:0}}
body{{background:var(--bg);color:var(--ink);font:14.5px/1.55 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Helvetica,Arial,sans-serif;-webkit-font-smoothing:antialiased}}
.top{{background:linear-gradient(100deg,var(--brand1),var(--brand2));color:#fff;padding:22px 0}}
.top .wrap{{display:flex;align-items:center;gap:14px}}
.logo{{width:38px;height:38px;border-radius:10px;background:rgba(255,255,255,.16);display:grid;place-items:center;font-weight:800;font-size:20px;flex:0 0 auto}}
.top h1{{font-size:19px;margin:0;font-weight:700;letter-spacing:.2px}}
.top .meta{{margin-left:auto;text-align:right;font-size:12px;opacity:.9}}
.wrap{{max-width:1160px;margin:0 auto;padding:0 22px}}
.body{{padding:26px 0 70px}}
.kpis{{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:14px;margin-bottom:26px}}
.kpi{{background:var(--card);border:1px solid var(--line);border-radius:14px;padding:16px 18px;box-shadow:var(--shadow)}}
.kval{{font-size:28px;font-weight:750;letter-spacing:-.5px}}
.klabel{{color:var(--muted);font-size:11.5px;text-transform:uppercase;letter-spacing:.05em;margin-top:3px;font-weight:600}}
.ksub{{color:var(--muted);font-size:11px;margin-top:3px}}
.grid2{{display:grid;grid-template-columns:1fr 1fr;gap:18px;margin-bottom:20px}}
@media(max-width:800px){{.grid2{{grid-template-columns:1fr}}}}
.panel{{background:var(--card);border:1px solid var(--line);border-radius:14px;padding:18px 20px;box-shadow:var(--shadow)}}
.panel h2{{font-size:12px;text-transform:uppercase;letter-spacing:.05em;color:var(--muted);margin:0 0 14px;font-weight:700}}
.brow{{display:grid;grid-template-columns:150px 1fr 32px;gap:12px;align-items:center;margin:9px 0}}
.blabel{{font-size:13px;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}}
.btrack{{background:var(--line);border-radius:7px;height:9px;overflow:hidden}}
.bfill{{height:100%;border-radius:7px;transition:width .3s}}
.bval{{text-align:right;font-variant-numeric:tabular-nums;color:var(--muted);font-size:13px}}
.tablewrap{{background:var(--card);border:1px solid var(--line);border-radius:14px;box-shadow:var(--shadow);overflow-x:auto;margin-top:20px}}
table{{width:100%;border-collapse:collapse;min-width:820px}}
th,td{{padding:12px 14px;text-align:left;border-bottom:1px solid var(--line);vertical-align:top}}
th{{font-size:11px;text-transform:uppercase;letter-spacing:.05em;color:var(--muted);font-weight:700;position:sticky;top:0;background:var(--card)}}
tr:last-child td{{border-bottom:none}}
tbody tr:hover{{background:color-mix(in srgb,var(--accent) 5%,transparent)}}
.num{{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}}
.btitle{{font-weight:650}}
.bsub{{color:var(--muted);font-size:12px;margin-top:2px}}
.pill{{display:inline-block;padding:2px 10px;border-radius:999px;font-size:11.5px;font-weight:650;color:#fff;background:var(--c);white-space:nowrap}}
.muted{{color:var(--muted)}}
.foot{{color:var(--muted);font-size:12px;margin-top:18px;text-align:center}}
h2.sec{{font-size:13px;text-transform:uppercase;letter-spacing:.05em;color:var(--muted);margin:6px 2px 12px;font-weight:700}}
</style></head>
<body>
<div class="top"><div class="wrap">
  <div class="logo">G</div>
  <div><h1>Genie &nbsp;·&nbsp; Production Dashboard</h1>
  <div style="font-size:12px;opacity:.9">Books produced by Genie — Claude Code &amp; Codex</div></div>
  <div class="meta">Updated {now}<br>{total} books tracked</div>
</div></div>
<div class="wrap body">
  <div class="kpis">{kpis}</div>
  <h2 class="sec">Overview</h2>
  <div class="grid2">
    <div class="panel"><h2>By status</h2>{status_bars}</div>
    <div class="panel"><h2>By producing agent</h2>{agent_bars}</div>
  </div>
  <div class="grid2">
    <div class="panel"><h2>By niche</h2>{niche_bars}</div>
    <div class="panel"><h2>By month</h2>{month_bars}</div>
  </div>
  <h2 class="sec">Books</h2>
  <div class="tablewrap"><table>
    <thead><tr><th>Title</th><th>Author</th><th>Niche</th><th>Agent</th><th>Status</th>
      <th>Trademark</th><th class="num">Cover</th><th class="num">Words</th><th class="num">Date</th><th class="num">Last activity</th></tr></thead>
    <tbody>{table}</tbody></table></div>
  <div class="foot">Only genuine Genie books (production record + catalog) · Agent from the failover audit log · Source: {src}</div>
</div></body></html>"""

if __name__ == "__main__":
    cat_path, books = build()
    with open(OUT, "w", encoding="utf-8") as f:
        f.write(render(cat_path, books))
    print(f"Wrote {OUT}")
    print(f"Catalog: {cat_path}")
    print(f"Genie books: {len(books)}")
    for b in books:
        print(f"  - {b['title'][:40]:40} | {b['agent']:11} | {b['status_norm']:12} | {b['niche']}")
